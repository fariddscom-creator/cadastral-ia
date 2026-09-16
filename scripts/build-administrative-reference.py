import json
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'grappes_echantillon.xlsx'
TARGET = ROOT / 'data' / 'administrative-reference.json'


def code(value: object, width: int) -> str:
    return str(int(value)).zfill(width)


workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook.active
header = next(sheet.iter_rows(values_only=True))
columns = {name: index for index, name in enumerate(header)}
districts: dict[str, dict[str, str]] = {}

for row in sheet.iter_rows(values_only=True, min_row=2):
    wilaya = row[columns['Wilaya']]
    commune = row[columns['Commune']]
    district = row[columns['District']]
    commune_name = row[columns['Com_fr']]
    if None in (wilaya, commune, district) or not isinstance(commune_name, str) or not commune_name.strip():
        continue
    key = f'{code(wilaya, 2)}.{code(commune, 2)}.{code(district, 3)}'
    districts[key] = {'commune': commune_name.strip()}

TARGET.parent.mkdir(parents=True, exist_ok=True)
TARGET.write_text(json.dumps({'source': SOURCE.name, 'districts': districts}, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
